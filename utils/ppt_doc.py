import os
import re
import json
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.text import PP_ALIGN
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor
from datetime import datetime
from pathlib import Path
import logging
from utils.image_generator import get_image_path_from_description

logger = logging.getLogger(__name__)

def add_chart_to_slide(slide, chart_data):
    """
    Slaydga diagramma (Chart) qo'shish
    chart_data: {
        "type": "column", "pie", "line", "bar"
        "title": str,
        "categories": [str],
        "series": [{"name": str, "values": [float]}]
    }
    """
    try:
        if not chart_data:
            return

        chart_type_map = {
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE,
            "pie": XL_CHART_TYPE.PIE
        }
        
        c_type = chart_type_map.get(chart_data.get("type", "column"), XL_CHART_TYPE.COLUMN_CLUSTERED)
        
        chart_data_obj = CategoryChartData()
        chart_data_obj.categories = chart_data.get("categories", [])
        
        for serie in chart_data.get("series", []):
            chart_data_obj.add_series(serie.get("name", "Series"), serie.get("values", []))
            
        # Define chart position (Bottom Right or Right Half usually)
        x, y, cx, cy = PptInches(5), PptInches(1.5), PptInches(4.5), PptInches(5)
        
        chart = slide.shapes.add_chart(
            c_type, x, y, cx, cy, chart_data_obj
        ).chart
        
        if chart_data.get("title"):
            chart.has_title = True
            chart.chart_title.text_frame.text = chart_data.get("title")
            
        # Adjust text layout if chart is present
        adjust_text_for_layout(slide, layout_type="content_with_image", image_position="right_half")
        
    except Exception as e:
        logger.error(f"Error adding chart to slide: {e}")

def add_image_to_placeholder(slide, image_path):
    try:
        # First pass: try to find a Picture placeholder (type 18)
        for placeholder in slide.placeholders:
            pf = getattr(placeholder, "placeholder_format", None)
            if pf and getattr(pf, "type", None) == 18:  # 18 = PICTURE
                placeholder.insert_picture(image_path)
                return True
        
        # Second pass: if no picture placeholder, DO NOT insert blindly.
        # We want to avoid covering text.
        # But if we must insert and no placeholder exists, we rely on add_image_to_slide's positioning.
        
    except Exception as e:
        logger.error(f"Error inserting image into placeholder: {e}")
    return False

def adjust_text_for_layout(slide, layout_type="content_with_image", image_position="right_half"):
    """
    Slayd tahrirlash: Matn va rasm/diagramma o'rtasidagi muvozanatni saqlash.
    """
    try:
        for shape in slide.placeholders:
            if not getattr(shape, "has_text_frame", False):
                continue
                
            # Sarlavhani (Title) odatda yuqorida qoldiramiz
            if getattr(shape.placeholder_format, "type", None) == 1: # 1 is TITLE
                shape.left = PptInches(0.5)
                shape.top = PptInches(0.3)
                shape.width = PptInches(9)
                shape.height = PptInches(1)
                continue

            # Asosiy matn (Body)
            if layout_type == "content_with_image":
                if image_position == "right_half":
                    shape.left = PptInches(0.5)
                    shape.top = PptInches(1.5)
                    shape.width = PptInches(4.5)
                    shape.height = PptInches(5)
                elif image_position == "left_half":
                    shape.left = PptInches(5)
                    shape.top = PptInches(1.5)
                    shape.width = PptInches(4.5)
                    shape.height = PptInches(5)
            elif layout_type == "comparison_slide":
                # Comparison slide handled differently (two text boxes usually)
                pass
            elif layout_type == "quote_slide":
                shape.left = PptInches(1)
                shape.top = PptInches(2.5)
                shape.width = PptInches(8)
                shape.height = PptInches(3)
                # Center alignment for quotes
                for paragraph in shape.text_frame.paragraphs:
                    paragraph.alignment = PP_ALIGN.CENTER
                    
    except Exception as e:
        logger.error(f"Error adjusting text for layout: {e}")

def fill_slide_text(slide, content_list):
    """Slaydga matn qo'shish va shriftni avtomatik sozlash"""
    try:
        tf = None
        if len(slide.placeholders) > 1:
            body_shape = slide.placeholders[1]
            if body_shape.has_text_frame:
                tf = body_shape.text_frame
        
        if not tf:
            # Fallback to creating a textbox if no placeholder
            txBox = slide.shapes.add_textbox(PptInches(1), PptInches(2), PptInches(8), PptInches(4))
            tf = txBox.text_frame

        tf.clear()
        tf.word_wrap = True
        
        # Determine font size based on content length
        total_chars = sum(len(str(item)) for item in content_list)
        if total_chars > 800:
            font_size = PptPt(14)
        elif total_chars > 500:
            font_size = PptPt(16)
        else:
            font_size = PptPt(20)

        for item in content_list:
            p = tf.add_paragraph()
            p.text = str(item)
            p.level = 0
            p.space_after = PptPt(10)
            p.font.size = font_size
            # Har bir paragraf uchun alignmentni tekshirish (agar kerak bo'lsa)
            # p.alignment = PP_ALIGN.LEFT # Standart chapdan
            
        return tf
    except Exception as e:
        logger.error(f"Error filling slide text: {e}")
        return None

def add_image_to_slide(slide, image_path, position="right"):
    """
    Slaydga rasm qo'shish
    position: "right", "left", "center", "full_background", "bottom_right", "top_left", "top_right", "left_half", "right_half"
    """
    try:
        if not os.path.exists(image_path):
            return

        from PIL import Image as PILImage
        
        # Slayd o'lchamlari (odatiy 10x7.5 dyuym)
        SLIDE_WIDTH = PptInches(10)
        SLIDE_HEIGHT = PptInches(7.5)

        # Rasm o'lchamlarini aniqlash
        with PILImage.open(image_path) as img:
            img_w, img_h = img.size
            aspect_ratio = img_w / img_h

        if position == "full_background":
            pic = slide.shapes.add_picture(image_path, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
            try:
                slide.shapes._spTree.insert(2, pic._element)
            except:
                pass
            return

        # "Object-fit: contain" mantiqini qo'llash
        def get_contain_dimensions(max_w, max_h, aspect):
            if (max_w / max_h) > aspect:
                # Balandlik cheklovchi faktor
                h = max_h
                w = h * aspect
            else:
                # Kenglik cheklovchi faktor
                w = max_w
                h = w / aspect
            return w, h

        if position == "center":
            max_w, max_h = PptInches(8), PptInches(5)
            w, h = get_contain_dimensions(max_w, max_h, aspect_ratio)
            left = (SLIDE_WIDTH - w) / 2
            top = (SLIDE_HEIGHT - h) / 2
            slide.shapes.add_picture(image_path, int(left), int(top), width=int(w), height=int(h))
            return

        elif position == "right_half":
            max_w, max_h = PptInches(4.5), PptInches(5)
            w, h = get_contain_dimensions(max_w, max_h, aspect_ratio)
            left = PptInches(5.25) + (max_w - w) / 2
            top = PptInches(1.5) + (max_h - h) / 2
            slide.shapes.add_picture(image_path, int(left), int(top), width=int(w), height=int(h))
            adjust_text_for_layout(slide, layout_type="content_with_image", image_position="right_half")
            return

        elif position == "left_half":
            max_w, max_h = PptInches(4.5), PptInches(5)
            w, h = get_contain_dimensions(max_w, max_h, aspect_ratio)
            left = PptInches(0.25) + (max_w - w) / 2
            top = PptInches(1.5) + (max_h - h) / 2
            slide.shapes.add_picture(image_path, int(left), int(top), width=int(w), height=int(h))
            adjust_text_for_layout(slide, layout_type="content_with_image", image_position="left_half")
            return

        elif position == "top_left":
            max_w, max_h = PptInches(3), PptInches(2.5)
            w, h = get_contain_dimensions(max_w, max_h, aspect_ratio)
            slide.shapes.add_picture(image_path, PptInches(0.5), PptInches(1.5), width=int(w), height=int(h))
            
        elif position == "top_right":
            max_w, max_h = PptInches(3), PptInches(2.5)
            w, h = get_contain_dimensions(max_w, max_h, aspect_ratio)
            slide.shapes.add_picture(image_path, PptInches(6.5), PptInches(1.5), width=int(w), height=int(h))

        else: # "right" or default
            max_w = PptInches(4)
            slide.shapes.add_picture(image_path, PptInches(5.5), PptInches(2), width=max_w)

    except Exception as e:
        logger.error(f"Error adding image to slide: {e}")

async def create_pptx_file(content_json: str, tema: str, muallif: str, user_id: int = None, shablon_id: int = 1, images: list = None) -> str:
    """PPTX fayl yaratish (JSON format asosida)"""
    try:
        # Template path checking
        template_path = Path(f"templates/design_{shablon_id}.pptx")
        if template_path.exists():
            prs = Presentation(str(template_path))
        else:
            prs = Presentation()
            
        # Parse JSON content
        try:
            # Clean up potential markdown code blocks
            clean_content = content_json.strip()
            if clean_content.startswith("```json"):
                clean_content = clean_content[7:]
            elif clean_content.startswith("```"):
                 clean_content = clean_content[3:]
            if clean_content.endswith("```"):
                clean_content = clean_content[:-3]
            
            slides_data = json.loads(clean_content)
        except json.JSONDecodeError as e:
            logger.error(f"JSON Decode Error: {e}")
            logger.error(f"Raw content: {content_json[:200]}...")
            return None
        
        # Sort slides by slide_number just in case
        slides_data.sort(key=lambda x: x.get("slide_number", 0))

        # Image distribution logic
        # We will map available images to slides that need them
        available_images = images.copy() if images else []
        
        for slide_data in slides_data:
            slide_num = slide_data.get("slide_number")
            slide_type = slide_data.get("slide_type", "content_slide")
            title_text = slide_data.get("title", "")
            subtitle_text = slide_data.get("subtitle", "")
            content_list = slide_data.get("content", [])
            speaker_notes = slide_data.get("speaker_notes", "")
            image_suggestion = slide_data.get("image_suggestion", {})
            
            # Select Layout
            # 0: Title, 1: Title+Content, 6: Title+Picture (in some templates), 8: Picture w/ Caption
            if slide_type == "title_slide":
                layout_index = 0
            elif slide_type == "image_slide":
                layout_index = 8 # Picture with Caption often good for image focus
            elif slide_type == "content_with_image":
                layout_index = 1 # We'll manually place image
            elif slide_type == "conclusion_slide":
                layout_index = 0 # Use Title slide style for impact
            else:
                layout_index = 1 # Standard Content
            
            # Fallback if layout doesn't exist
            if layout_index >= len(prs.slide_layouts):
                layout_index = 1 
                
            slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
            
            # Add Title
            if slide.shapes.title:
                slide.shapes.title.text = title_text
            
            # Add Subtitle (for Title Slide)
            if slide_type == "title_slide":
                try:
                    if len(slide.placeholders) > 1:
                        subtitle = slide.placeholders[1]
                        subtitle.text = f"{subtitle_text}\nTayyorladi: {muallif}\nMavzu: {tema}"
                except:
                    pass

            # Add Content (Bullet points / Paragraphs)
            if content_list and slide_type not in ["title_slide", "conclusion_slide"]:
                fill_slide_text(slide, content_list)

            # Adjust Layout based on slide_type
            if slide_type == "quote_slide":
                adjust_text_for_layout(slide, layout_type="quote_slide")
            elif slide_type == "comparison_slide":
                # For comparison, we might need a specific layout or custom textboxes
                pass

            # FORCE EMPTY SPEAKER NOTES
            if slide.has_notes_slide:
                try:
                    slide.notes_slide.notes_text_frame.text = ""
                except:
                    pass

            # Add Chart if available
            chart_data = slide_data.get("chart_data")
            if chart_data:
                add_chart_to_slide(slide, chart_data)
                
            # Add Image
            needs_image = slide_type in ["image_slide", "content_with_image", "quote_slide"] or image_suggestion
            
            # If chart exists, we usually don't want image unless specified
            if chart_data and not image_suggestion:
                needs_image = False
            
            img_path = None
            if needs_image:
                if available_images:
                    img_path = available_images.pop(0) 
                elif image_suggestion and image_suggestion.get("description"):
                    try:
                        img_path = await get_image_path_from_description(
                            image_suggestion.get("description"),
                            image_suggestion.get("style", "modern, professional")
                        )
                    except Exception as e:
                        logger.error(f"Error generating image: {e}")

            if img_path:
                position = image_suggestion.get("position", "right_half") if image_suggestion else "right_half"
                
                if slide_type == "image_slide":
                    position = "center"
                elif slide_type == "quote_slide":
                    position = "full_background"
                
                inserted = add_image_to_placeholder(slide, img_path)
                if not inserted:
                    add_image_to_slide(slide, img_path, position=position)
                    # For custom images, we might need to re-adjust text
                    if position in ["right_half", "left_half"]:
                        adjust_text_for_layout(slide, layout_type="content_with_image", image_position=position)

        # Save File
        clean_tema = "".join([c for c in tema if c.isalpha() or c.isdigit() or c in (' ', '-', '_')]).strip()
        filename = f"taqdimot_{clean_tema[:20]}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
        
        # Folder structure based on user_id
        if user_id:
            base_path = Path("documents") / str(user_id) / "taqdimotlar"
        else:
            base_path = Path("taqdimotlar")
            
        base_path.mkdir(parents=True, exist_ok=True)
        
        full_path = base_path / filename
        prs.save(str(full_path))
        
        # Save metadata info file if user_id is present
        if user_id:
            try:
                info_filename = f"taqdimot_{clean_tema[:20]}_info.txt"
                info_path = base_path / info_filename
                with open(info_path, "w", encoding="utf-8") as f:
                    f.write(f"Mavzu: {tema}\n")
                    f.write(f"Muallif: {muallif}\n")
                    f.write(f"Shablon ID: {shablon_id}\n")
                    f.write(f"Vaqt: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                    if images:
                        f.write(f"Rasmlar soni: {len(images)}\n")
            except Exception as e:
                logger.error(f"Error saving ppt metadata info: {e}")
        
        return str(full_path)
        
    except Exception as e:
        logger.error(f"❌ PPTX fayl yaratishda xatolik: {str(e)}", exc_info=True)
        return None
